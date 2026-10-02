"""The workbench routes (`/workbench/...`), on a toy synthetic population."""

from __future__ import annotations

import io
import shutil
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from poltika.api import microsim, workbench, workbench_registry
from poltika.api.main import app
from poltika.scenarios.microsim.runner import COMPONENTS
from poltika.scenarios.microsim.synthetic import generate
from poltika.semantics.population import PopulationDataset
from tests.unit.scenarios.microsim.test_synthetic import _spec

client = TestClient(app)
THRESHOLD_PARAM = "family_tax_credit_abatement_threshold"


def _reform(value: str = "50000", reform_id: str = "threshold") -> dict[str, Any]:
    return {
        "reform_id": reform_id,
        "overrides": [
            {
                "policy_id": "working_for_families",
                "parameter_id": THRESHOLD_PARAM,
                "operation": "set",
                "value": value,
            }
        ],
    }


def _results_body(label: str = "Option A") -> dict[str, Any]:
    return {"scenario": {"name": "Test"}, "options": [{"label": label, "reform": _reform()}]}


@pytest.fixture
def configured(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    population, _ = generate(_spec(households=40))
    path = tmp_path / "population.json"
    path.write_text(population.model_dump_json())
    monkeypatch.setenv("POLTIKA_POPULATION_PATH", str(path))
    return path


def _table(payload: dict[str, Any], table_id: str) -> dict[str, Any]:
    for tab in payload["tabs"]:
        for block in tab["blocks"]:
            if block["id"] == table_id:
                return block
    raise AssertionError(f"no table {table_id}")


# --- registry ---------------------------------------------------------------


def test_registry_lists_instruments_parameters_and_sources() -> None:
    body = client.get("/workbench/registry").json()
    by_id = {i["id"]: i for i in body["instruments"]}
    assert {"working_for_families", "accommodation_supplement"} <= set(by_id)
    wff = by_id["working_for_families"]
    assert wff["runnable"] is True and wff["valid"] is True
    threshold = next(p for p in wff["parameters"] if p["id"] == THRESHOLD_PARAM)
    assert threshold["editable"] is True
    assert threshold["current"]["outcome"] == "resolved"
    assert threshold["current"]["source_id"] in {s["id"] for s in wff["sources"]}
    source = next(s for s in wff["sources"] if s["id"] == threshold["current"]["source_id"])
    assert source["url"].startswith("https://") and source["title"]
    assert wff["decisions"] and wff["decisions"][0]["id"] == "DEC-001"
    assert [t["id"] for t in body["tabs"] if t["status"] == "active"] == [
        "fiscal",
        "households",
        "assumptions",
    ]
    assert body["behavioural"]["available"] is False
    assert body["max_options"] == 2


def test_registry_never_returns_a_filesystem_path() -> None:
    text = client.get("/workbench/registry").text
    assert str(workbench_registry.specs_root()) not in text


def test_a_new_instrument_appears_with_no_code_change(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The registry is driven by the spec directory: a third instrument is
    listed, read-only, with its parameters and sources, and nothing else
    changes."""
    root = tmp_path / "specs"
    shutil.copytree(workbench_registry.specs_root(), root)
    dummy = root / "dummy_instrument"
    shutil.copytree(root / "accommodation_supplement", dummy)
    policy = (dummy / "policy.yaml").read_text().replace(
        "policy_id: accommodation_supplement", "policy_id: dummy_instrument"
    )
    (dummy / "policy.yaml").write_text(policy)
    monkeypatch.setattr(workbench_registry, "specs_root", lambda: root)

    instruments = client.get("/workbench/registry").json()["instruments"]
    added = next(i for i in instruments if i["id"] == "dummy_instrument")
    assert added["valid"] is True and added["runnable"] is False
    assert added["parameters"] and all(p["editable"] is False for p in added["parameters"])
    assert added["sources"]
    assert {i["id"] for i in instruments} >= {
        "working_for_families",
        "accommodation_supplement",
        "dummy_instrument",
    }


def test_an_unparseable_instrument_is_reported_and_hides_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "specs"
    shutil.copytree(workbench_registry.specs_root(), root)
    (root / "broken").mkdir()
    (root / "broken" / "policy.yaml").write_text("policy_id: [unclosed")
    monkeypatch.setattr(workbench_registry, "specs_root", lambda: root)

    instruments = client.get("/workbench/registry").json()["instruments"]
    broken = next(i for i in instruments if i["id"] == "broken")
    assert broken["valid"] is False and broken["error"]
    assert next(i for i in instruments if i["id"] == "working_for_families")["valid"] is True


# --- check ------------------------------------------------------------------


def test_check_reports_the_resulting_value() -> None:
    body = client.post("/workbench/check", json={"reform": _reform("50000")}).json()
    assert body["ok"] is True and body["problems"] == []
    change = body["changes"][0]
    assert change["resulting"] == "50000" and change["baseline"] not in (None, "50000")


def test_check_add_and_multiply_are_computed_by_the_engine() -> None:
    reform = {
        "reform_id": "m",
        "overrides": [
            {
                "policy_id": "working_for_families",
                "parameter_id": THRESHOLD_PARAM,
                "operation": "multiply",
                "value": "2",
            }
        ],
    }
    body = client.post("/workbench/check", json={"reform": reform}).json()
    assert Decimal(body["changes"][0]["resulting"]) == Decimal(body["changes"][0]["baseline"]) * 2


def test_check_names_a_refused_override() -> None:
    reform = {
        "reform_id": "bad",
        "overrides": [
            {
                "policy_id": "working_for_families",
                "parameter_id": "family_tax_credit_abatement_rate",
                "operation": "set",
                "value": "1.5",
            },
            {
                "policy_id": "working_for_families",
                "parameter_id": "no_such_parameter",
                "operation": "set",
                "value": "1",
            },
            {
                "policy_id": "income_tax",
                "parameter_id": "anything",
                "operation": "set",
                "value": "1",
            },
        ],
    }
    body = client.post("/workbench/check", json={"reform": reform}).json()
    assert body["ok"] is False and body["changes"] == []
    assert {p["parameter_id"] for p in body["problems"]} == {
        "family_tax_credit_abatement_rate",
        "no_such_parameter",
        "anything",
    }
    assert all(p["message"] for p in body["problems"])


def test_check_rejects_an_unknown_field() -> None:
    assert client.post("/workbench/check", json={"reform": _reform(), "x": 1}).status_code == 422


# --- results -----------------------------------------------------------------


def test_results_need_a_population(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("POLTIKA_POPULATION_PATH", raising=False)
    assert client.post("/workbench/results", json=_results_body()).status_code == 503


def test_workbench_population_run_honours_shared_admission(configured: Path) -> None:
    with microsim._one_at_a_time():
        response = client.post("/workbench/results", json=_results_body())
    assert response.status_code == 429
    assert response.headers["Retry-After"] == "30"


def test_results_payload(configured: Path) -> None:
    response = client.post("/workbench/results", json=_results_body())
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["synthetic"] is True and "not an estimate" in payload["population_label"]
    run = payload["runs"][0]
    assert len(run["run_id"]) == 64 and run["label"] == "Option A"
    assert [t["id"] for t in payload["tabs"]][:4] == [
        "fiscal",
        "households",
        "work_incentives",
        "assumptions",
    ]
    assert {t["id"]: t["status"] for t in payload["tabs"]}["cba"] == "deferred"
    horizon = next(b for b in _tab(payload, "fiscal")["blocks"] if b["type"] == "horizon")
    assert [y["modelled"] for y in horizon["years"]] == [True, False, False, False]


def _tab(payload: dict[str, Any], tab_id: str) -> dict[str, Any]:
    return next(t for t in payload["tabs"] if t["id"] == tab_id)


def test_fiscal_table_has_a_row_per_component_and_no_total(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    table = _table(payload, "fiscal_components")
    assert [row[0] for row in table["rows"]] == list(COMPONENTS)
    assert not any("total" in str(row[0]).lower() for row in table["rows"])


def test_a_component_that_is_not_calculated_is_never_zero(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    rows = {r[0]: r for r in _table(payload, "fiscal_components")["rows"]}
    # The 2026-27 in-work tax credit is gated, so it is never calculated.
    assert rows["in_work_tax_credit"][2:] == ["not calculated"] * 3
    examples = _table(payload, "worked_examples")["rows"]
    iwtc = [r for r in examples if r[1] == "in_work_tax_credit"]
    outcomes = {"needs_review", "unsupported"}  # never a number, never zero
    assert iwtc and all(cell in outcomes for r in iwtc for cell in (r[3], r[4], r[5]))


def test_headline_matches_the_engine_fiscal_totals(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    engine = client.post(
        "/microsim/runs", json={"reform": _reform(), "component_set": ["family_tax_credit"]}
    ).json()
    assert engine["run_id"] == payload["runs"][0]["run_id"]
    ftc = next(h for h in payload["headline"] if h["measure_id"] == "family_tax_credit")
    exact = Decimal(engine["fiscal"]["family_tax_credit"]["delta_total"])
    assert Decimal(ftc["change_total"]) == exact.quantize(Decimal("0.01"))
    assert ftc["threshold"] == "52"


def test_a_worked_example_equals_the_household_route(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    family = client.get("/workbench/registry").json()["family_types"][1]
    request = next(
        e["request"] for e in client.get("/analyst/examples").json() if e["id"] == family["id"]
    )
    direct = client.post("/household/calculate", json=request).json()
    rows = [r for r in _table(payload, "worked_examples")["rows"] if r[0] == family["description"]]
    ftc = next(r for r in rows if r[1] == "family_tax_credit")
    assert ftc[3] == direct["wff"]["family_tax_credit"]["value"]


def test_two_options_are_compared_to_one_status_quo(configured: Path) -> None:
    body = {
        "scenario": {"name": "Two", "version": "2", "status": "final"},
        "options": [
            {"label": "A", "reform": _reform("50000", "a")},
            {"label": "B", "reform": _reform("60000", "b")},
        ],
    }
    payload = client.post("/workbench/results", json=body).json()
    assert [r["label"] for r in payload["runs"]] == ["A", "B"]
    assert payload["runs"][0]["run_id"] != payload["runs"][1]["run_id"]
    table = _table(payload, "fiscal_components")
    assert [c["label"] for c in table["columns"]][2:] == [
        "Status quo",
        "A",
        "A: change",
        "B",
        "B: change",
    ]
    # One status quo, and it is each run's own baseline.
    for label, threshold in (("a", "50000"), ("b", "60000")):
        engine = client.post(
            "/microsim/runs",
            json={"reform": _reform(threshold, label), "component_set": ["family_tax_credit"]},
        ).json()
        ftc = next(r for r in table["rows"] if r[0] == "family_tax_credit")
        baseline = Decimal(engine["fiscal"]["family_tax_credit"]["baseline_total"])
        assert Decimal(ftc[2]) == baseline.quantize(Decimal("0.01"))


def test_option_labels_must_be_unique_and_not_the_status_quo(configured: Path) -> None:
    option = {"label": "A", "reform": _reform()}
    assert (
        client.post("/workbench/results", json={"options": [option, option]}).status_code == 422
    )
    named = {"label": "Status quo", "reform": _reform()}
    assert client.post("/workbench/results", json={"options": [named]}).status_code == 422


def test_at_most_two_options(configured: Path) -> None:
    options = [{"label": str(i), "reform": _reform(str(50000 + i), str(i))} for i in range(3)]
    assert client.post("/workbench/results", json={"options": options}).status_code == 422


def test_an_invalid_reform_is_a_422_not_a_500(configured: Path) -> None:
    reform = _reform("-5")
    body = {"options": [{"label": "A", "reform": reform}]}
    assert client.post("/workbench/results", json=body).status_code == 422


def test_assumptions_carry_run_identity_sources_and_limitations(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    tab = _tab(payload, "assumptions")
    identity = {r[0]: r[1] for r in _table(payload, "run_identity")["rows"]}
    assert identity["Run ID: Option A"] == payload["runs"][0]["run_id"]
    assert identity["Law baseline"] == payload["law_baseline"]
    used = _table(payload, "parameters_used")["rows"]
    assert used and all(row[8] and row[10] for row in used if row[7] == "approved")
    changed = _table(payload, "parameters_changed")["rows"]
    assert [r[2] for r in changed] == [THRESHOLD_PARAM] and changed[0][7].startswith("scenario")
    assert _table(payload, "decisions")["rows"]
    limitations = next(b for b in tab["blocks"] if b["id"] == "limitations")["items"]
    assert any("Static costing" in s for s in limitations)
    assert any("not an official poverty measure" in s.lower() for s in limitations)


# --- export ------------------------------------------------------------------


def _export_body(payload: dict[str, Any], label: str = "Option A") -> dict[str, Any]:
    return {
        "scenario": {"name": "Test"},
        "options": [{"label": label, "run_id": payload["runs"][0]["run_id"]}],
    }


def test_excel_totals_match_the_screen(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    response = client.post("/workbench/export.xlsx", json=_export_body(payload))
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Summary", "Fiscal", "Households", "Assumptions & provenance"]
    run_id = payload["runs"][0]["run_id"]
    for ws in wb.worksheets:
        assert run_id in str(ws["A2"].value)

    table = _table(payload, "fiscal_components")
    sheet = wb["Fiscal"]
    rows = list(sheet.iter_rows(values_only=True))
    header_at = next(
        i for i, r in enumerate(rows) if r[0] == "Component" and r[2] == "Status quo"
    )
    for offset, expected in enumerate(table["rows"], start=1):
        actual = rows[header_at + offset]
        assert actual[0] == expected[0]
        for got, want in zip(actual[2:5], expected[2:5], strict=True):
            if want == "not calculated":
                assert got == want
            else:
                assert isinstance(got, (int, float, Decimal))
                assert Decimal(str(got)) == Decimal(want)


def test_excel_headline_matches_the_screen(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    wb = load_workbook(
        io.BytesIO(client.post("/workbench/export.xlsx", json=_export_body(payload)).content)
    )
    rows = list(wb["Summary"].iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Measure")
    for offset, h in enumerate(payload["headline"], start=1):
        row = rows[start + offset]
        assert row[0] == h["measure"] and row[1] == h["option"]
        assert Decimal(str(row[3])) == Decimal(h["change_total"])


def test_csv_per_table(configured: Path) -> None:
    payload = client.post("/workbench/results", json=_results_body()).json()
    response = client.post("/workbench/export/fiscal_components.csv", json=_export_body(payload))
    assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
    lines = response.text.splitlines()
    assert payload["runs"][0]["run_id"] in lines[0]
    assert lines[1].startswith("Component,Unit,Status quo")
    assert len(lines) == 2 + len(COMPONENTS)
    missing = client.post("/workbench/export/nope.csv", json=_export_body(payload))
    assert missing.status_code == 404


def test_formula_like_labels_are_written_as_text(configured: Path) -> None:
    label = "=HYPERLINK(1)"
    payload = client.post("/workbench/results", json=_results_body(label)).json()
    body = _export_body(payload, label)
    wb = load_workbook(io.BytesIO(client.post("/workbench/export.xlsx", json=body).content))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value == label]
    assert cells and all(c.data_type == "s" for c in cells)
    csv_text = client.post("/workbench/export/fiscal_coverage.csv", json=body).text
    assert "'=HYPERLINK(1)" in csv_text
    header = client.post("/workbench/export/fiscal_components.csv", json=body).text.splitlines()[1]
    assert "'=HYPERLINK(1): change" in header and ",=HYPERLINK" not in header


def test_export_of_an_unknown_run_is_404(configured: Path) -> None:
    body = {"options": [{"label": "A", "run_id": "0" * 64}]}
    assert client.post("/workbench/export.xlsx", json=body).status_code == 404
    assert client.post("/workbench/export.xlsx", json={"options": [{"label": "A", "run_id": "x"}]}).status_code == 422


# --- decisions -----------------------------------------------------------------


def test_decision_text_and_path_safety() -> None:
    ok = client.get("/workbench/decisions/accommodation_supplement/DEC-003")
    assert ok.status_code == 200 and ok.text.startswith("# DEC-003")
    assert client.get("/workbench/decisions/accommodation_supplement/DEC-999").status_code == 404
    assert client.get("/workbench/decisions/..%2F..%2Fetc/DEC-001").status_code in (404, 422)
    assert client.get("/workbench/decisions/accommodation_supplement/..").status_code in (404, 422)


# --- preflight probes ------------------------------------------------------------


def test_a_broken_instrument_reveals_no_server_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "specs"
    shutil.copytree(workbench_registry.specs_root(), root)
    (root / "nosrc").mkdir()
    shutil.copy(root / "accommodation_supplement" / "policy.yaml", root / "nosrc" / "policy.yaml")
    monkeypatch.setattr(workbench_registry, "specs_root", lambda: root)

    response = client.get("/workbench/registry")
    broken = next(i for i in response.json()["instruments"] if i["id"] == "nosrc")
    assert broken["valid"] is False and broken["error"]
    assert str(tmp_path) not in response.text


def test_registry_as_at_outside_the_reference_period_says_unsupported_not_zero() -> None:
    body = client.get("/workbench/registry", params={"as_at": "2030-01-01"}).json()
    wff = next(i for i in body["instruments"] if i["id"] == "working_for_families")
    assert all(p["current"]["outcome"] == "unsupported" for p in wff["parameters"])
    assert all(p["current"]["value"] is None for p in wff["parameters"])
    assert client.get("/workbench/registry", params={"as_at": "nope"}).status_code == 422


def test_the_registry_title_drops_the_specification_suffix() -> None:
    body = client.get("/workbench/registry").json()
    assert {i["title"] for i in body["instruments"] if i["valid"]} >= {
        "Working for Families",
        "Accommodation Supplement",
    }


def test_entirely_unresolved_headline_is_not_zero_in_screen_or_exports(configured: Path) -> None:
    population = PopulationDataset.model_validate_json(configured.read_text())
    for household in population.households:
        household.facts["household"]["wff"]["family_scheme_income"] = None
    configured.write_text(population.model_dump_json())
    payload = client.post("/workbench/results", json=_results_body()).json()
    ftc = next(h for h in payload["headline"] if h["measure_id"] == "family_tax_credit")
    assert ftc["change_total"] == "not calculated" and ftc["mean_change"] is None
    assert Decimal(ftc["not_calculated"]) == sum(h.weight for h in population.households)
    body = _export_body(payload)
    csv_text = client.post("/workbench/export/headline.csv", json=body).text
    assert "Family tax credit,Option A,NZD per tax year,not calculated," in csv_text
    wb = load_workbook(io.BytesIO(client.post("/workbench/export.xlsx", json=body).content))
    row = next(r for r in wb["Summary"].iter_rows(values_only=True) if r[0] == "Family tax credit")
    assert row[3] == "not calculated" and row[4] is None


def test_invalid_households_are_counted_in_workbench_coverage(configured: Path) -> None:
    population = PopulationDataset.model_validate_json(configured.read_text())
    # One record has no claimant and is excluded before running.
    invalid = population.households[0]
    changed = invalid.model_copy(update={
        "persons": tuple(p for p in invalid.persons if p.role.value != "claimant")
    })
    population = population.model_copy(update={"households": (changed, *population.households[1:])})
    configured.write_text(population.model_dump_json())
    payload = client.post("/workbench/results", json=_results_body()).json()
    coverage = _table(payload, "fiscal_coverage")
    invalid_at = next(i for i, c in enumerate(coverage["columns"]) if c["id"] == "invalid")
    share_at = next(i for i, c in enumerate(coverage["columns"]) if c["id"] == "share")
    total_weight = sum(h.weight for h in population.households)
    for row in coverage["rows"]:
        assert Decimal(row[invalid_at]) == invalid.weight
        expected_share = (Decimal(row[2]) * 100 / total_weight).quantize(Decimal("0.1"))
        assert Decimal(row[share_at]) == expected_share
    winners = _table(payload, "winners_losers_family_tax_credit")
    invalid_at = next(i for i, c in enumerate(winners["columns"]) if c["id"] == "invalid")
    assert sum(Decimal(row[invalid_at]) for row in winners["rows"]) == invalid.weight


def test_export_refuses_different_populations_with_equal_baseline_totals(configured: Path) -> None:
    first = client.post("/workbench/results", json=_results_body()).json()
    population = PopulationDataset.model_validate_json(configured.read_text())
    configured.write_text(population.model_copy(update={"population_id": "other"}).model_dump_json())
    second = client.post("/workbench/results", json=_results_body()).json()
    assert _table(first, "fiscal_components")["rows"] == _table(second, "fiscal_components")["rows"]
    options = [
        {"label": label, "run_id": payload["runs"][0]["run_id"]}
        for label, payload in (("A", first), ("B", second))
    ]
    for path in ("/workbench/export.xlsx", "/workbench/export/fiscal_components.csv"):
        response = client.post(path, json={"options": options})
        assert response.status_code == 409
        assert "same population" in response.json()["detail"]


def test_export_refuses_changed_policy_artifacts(
    configured: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "specs"
    shutil.copytree(workbench_registry.specs_root(), root)
    monkeypatch.setattr(workbench_registry, "specs_root", lambda: root)
    monkeypatch.setattr(workbench, "specs_root", lambda: root)
    for module in (microsim, workbench):
        monkeypatch.setattr(module, "WFF_POLICY_DIR", root / "working_for_families")
        monkeypatch.setattr(module, "AS_POLICY_DIR", root / "accommodation_supplement")
    payload = client.post("/workbench/results", json=_results_body()).json()
    specification = root / "working_for_families" / "specification.md"
    specification.write_text(specification.read_text() + "\nChanged after this run.\n")
    for path in ("/workbench/export.xlsx", "/workbench/export/worked_examples.csv"):
        response = client.post(path, json=_export_body(payload))
        assert response.status_code == 409
        assert "policy files have changed" in response.json()["detail"]
        assert str(root) not in response.text


def test_check_config_errors_use_generic_server_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "specs"
    shutil.copytree(workbench_registry.specs_root(), root)
    policy = root / "working_for_families" / "policy.yaml"
    policy.write_text(policy.read_text().replace("source_catalog: sources.yaml", "source_catalog: missing.yaml"))
    monkeypatch.setattr(workbench, "specs_root", lambda: root)
    response = client.post("/workbench/check", json={"reform": _reform()})
    assert response.status_code == 500
    assert response.json() == {
        "detail": "Internal server error: the deployed policy configuration is invalid."
    }
    assert str(root) not in response.text
