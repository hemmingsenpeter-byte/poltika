"""Excel and CSV writers for the workbench payload.

Both are written from the same blocks the page renders (`workbench_tables`),
so the totals match the screen. Numbers are written as numbers, never as
images. Every sheet starts with the scenario name and the run IDs.

Text that a spreadsheet could read as a formula (a user-chosen option label
starting with `=`, `+`, `-` or `@`) is written as text.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

_NUMBER_FORMATS = {"money": "#,##0.00#", "count": "#,##0", "share": "0.0"}
_FORMULA_START = ("=", "+", "-", "@")


def _decimal(value: Any) -> Decimal | None:
    if isinstance(value, str):
        try:
            return Decimal(value)
        except InvalidOperation:
            return None
    return None


def _header_lines(payload: dict[str, Any]) -> list[list[str]]:
    scenario = payload["scenario"]
    runs = "; ".join(f"{r['label']}: {r['run_id']}" for r in payload["runs"])
    return [
        [f"Scenario: {scenario['name']} (version {scenario['version'] or '-'}, {scenario['status']})"],
        [f"Run ID: {runs}"],
        [f"Law baseline: {payload['law_baseline']}"],
        [payload["population_label"]],
    ]


def _write_cell(
    ws: Worksheet, row: int, column: int, value: Any, kind: str = "text", bold: bool = False
) -> None:
    """Every write goes through here, so no string is ever read as a formula."""
    cell = ws.cell(row=row, column=column)
    number = _decimal(value) if kind in _NUMBER_FORMATS else None
    if number is not None:
        cell.value = number
        cell.number_format = _NUMBER_FORMATS[kind]
    else:
        cell.value = value
        if isinstance(value, str) and value.startswith(_FORMULA_START):
            cell.data_type = "s"
    if bold:
        cell.font = Font(bold=True)


def _write_block(ws: Worksheet, row: int, block: dict[str, Any]) -> int:
    _write_cell(ws, row, 1, block["title"], bold=True)
    row += 1
    if block["type"] == "table":
        for c, column in enumerate(block["columns"], start=1):
            _write_cell(ws, row, c, column["label"], bold=True)
        row += 1
        for values in block["rows"]:
            for c, (value, column) in enumerate(zip(values, block["columns"], strict=True), start=1):
                _write_cell(ws, row, c, value, column["kind"])
            row += 1
        for note in block["notes"]:
            _write_cell(ws, row, 1, note)
            row += 1
    elif block["type"] == "horizon":
        for year in block["years"]:
            _write_cell(ws, row, 1, year["label"])
            _write_cell(ws, row, 2, "modelled" if year["modelled"] else "not modelled")
            _write_cell(ws, row, 3, year["note"])
            row += 1
    else:
        for item in block["items"]:
            _write_cell(ws, row, 1, item)
            row += 1
    return row + 1


def _sheet(wb: Workbook, title: str, payload: dict[str, Any], first: bool) -> Worksheet:
    ws = wb.active if first else wb.create_sheet()
    assert ws is not None
    ws.title = title[:31]
    for r, line in enumerate(_header_lines(payload), start=1):
        _write_cell(ws, r, 1, line[0])
    return ws


def headline_block(payload: dict[str, Any]) -> dict[str, Any]:
    columns = [
        {"id": "measure", "label": "Measure", "kind": "text"},
        {"id": "option", "label": "Option", "kind": "text"},
        {"id": "unit", "label": "Unit", "kind": "text"},
        {"id": "change_total", "label": "Change in total entitlement", "kind": "money"},
        {"id": "mean_change", "label": "Mean change of those calculated", "kind": "money"},
        {"id": "better_off", "label": "Better off (households)", "kind": "count"},
        {"id": "worse_off", "label": "Worse off (households)", "kind": "count"},
        {"id": "no_material_change", "label": "No material change", "kind": "count"},
        {"id": "not_calculated", "label": "Not calculated", "kind": "count"},
    ]
    rows = [[h[c["id"]] for c in columns] for h in payload["headline"]]
    return {
        "type": "table",
        "id": "headline",
        "title": "Headline",
        "columns": columns,
        "rows": rows,
        "notes": [],
    }


def to_xlsx(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    summary = _sheet(wb, "Summary", payload, first=True)
    _write_block(summary, len(_header_lines(payload)) + 2, headline_block(payload))
    for tab in payload["tabs"]:
        if tab["status"] != "active":
            continue
        ws = _sheet(wb, tab["title"], payload, first=False)
        row = len(_header_lines(payload)) + 2
        for block in tab["blocks"]:
            row = _write_block(ws, row, block)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def all_tables(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    tables = {"headline": headline_block(payload)}
    for tab in payload["tabs"]:
        for block in tab["blocks"]:
            if block["type"] == "table":
                tables[block["id"]] = block
    return tables


def _csv_cell(value: Any, kind: str) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        if kind in _NUMBER_FORMATS and _decimal(value) is not None:
            return value
        return "'" + value
    return value


def table_csv(payload: dict[str, Any], table_id: str) -> str | None:
    table = all_tables(payload).get(table_id)
    if table is None:
        return None
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow([f"Run ID: {r['label']}: {r['run_id']}" for r in payload["runs"]])
    writer.writerow([_csv_cell(c["label"], "text") for c in table["columns"]])
    for values in table["rows"]:
        writer.writerow(
            [_csv_cell(v, c["kind"]) for v, c in zip(values, table["columns"], strict=True)]
        )
    return out.getvalue()
